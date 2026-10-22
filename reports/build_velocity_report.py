"""Build the Korean sensitivity report from saved evaluation records."""
from pathlib import Path
import json
import math
import statistics

from docx import Document
from docx.shared import Inches, Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / 'logs/evaluation/velocity_sensitivity_shared_2026-10-21_vector'
OLD = ROOT / 'logs/evaluation/velocity_sensitivity_2026-10-21_combined'
OUT = ROOT / 'reports/ADAPT_velocity_sensitivity_report_2026-10-21.docx'
rows = json.loads((DATA / 'summary.json').read_text())
meta = json.loads((DATA / 'metadata.json').read_text())
trials = json.loads((DATA / 'trials.json').read_text())
old_trials = json.loads((OLD / 'trials.json').read_text())
by = {(r['noise_mps'], r['delay_ms']): r for r in rows}
assert len(trials) == 750 and sum(t['fell'] for t in trials) == 0
for r in trials:
    assert math.isclose(r['vxy_rmse_mps'] ** 2,
                        r['vx_rmse_mps'] ** 2 + r['vy_rmse_mps'] ** 2, rel_tol=1e-12)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.top_margin, sec.bottom_margin = Mm(18), Mm(18)
sec.left_margin, sec.right_margin = Mm(20), Mm(20)
sec.footer_distance = Mm(9)
FONT = 'Noto Sans CJK KR'
for name, size in [('Normal', 10), ('Title', 23), ('Subtitle', 11),
                   ('Heading 1', 15), ('Heading 2', 11.5), ('Caption', 8.5)]:
    style = doc.styles[name]
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn('w:' + key), FONT)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = Pt(size * 1.5)
    style.font.italic = False
    if name.startswith('Heading'):
        style.font.bold = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.keep_with_next = True
doc.styles['Caption'].font.italic = False
doc.styles['Title'].paragraph_format.space_after = Pt(9)
for element in doc.styles.element.iter():
    if element.tag in (qn('w:pBdr'), qn('w:spacing')) and (
        element.tag == qn('w:pBdr') or element.getparent().tag == qn('w:rPr')
    ):
        element.getparent().remove(element)
doc.core_properties.title = 'ADAPT 선속도 입력 노이즈와 지연에 따른 보행 성능 평가'
doc.core_properties.subject = 'actor와 observer가 공유하는 선속도 오차의 민감도 평가'
doc.core_properties.author = 'Codex'
doc.core_properties.keywords = 'ADAPT, 선속도, RMSE, observer, noise, delay'

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run()
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
run._r.addnext(field)

def p(text, bold=False, style=None):
    para = doc.add_paragraph(style=style)
    run = para.add_run(text)
    run.bold = bold
    return para

new_page = False

def h(text):
    global new_page
    para = doc.add_heading(text, 1)
    if new_page:
        para.paragraph_format.page_break_before = True
        new_page = False
    return para

def caption(text):
    para = p(text, style='Caption')
    return para

def table(headers, records, widths, font_size=9):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Mm(width)
    pr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        elem = OxmlElement('w:' + edge)
        for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            elem.set(qn('w:' + key), value)
        borders.append(elem)
    pr.append(borders)
    for rowidx, values in enumerate([headers] + records):
        row = t.rows[0] if rowidx == 0 else t.add_row()
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if rowidx == 0:
            trpr.append(OxmlElement('w:tblHeader'))
        for ci, (cell, value) in enumerate(zip(row.cells, values)):
            cell.width = Mm(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for edge in ('top', 'left', 'bottom', 'right'):
                el = OxmlElement('w:' + edge)
                el.set(qn('w:w'), '80')
                el.set(qn('w:type'), 'dxa')
                margins.append(el)
            tcpr.append(margins)
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'E4EDF4' if rowidx == 0 else ('F7F9FB' if rowidx % 2 else 'FFFFFF'))
            tcpr.append(shade)
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT if len(headers) == 2 and ci == 1 else WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = Pt(font_size * 1.5)
            run = para.add_run(str(value))
            run.font.size = Pt(font_size)
            run.bold = rowidx == 0
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(3)
    spacer.add_run().font.size = Pt(1)
    return t

def page():
    global new_page
    new_page = True

def equation(text):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.line_spacing = 1.0
    def node(tag, *children, text=None):
        el = OxmlElement('m:' + tag)
        for child in children:
            el.append(child)
        if text is not None:
            el.text = text
        return el
    def mr(text):
        return node('r', node('t', text=text))
    def sub(base, suffix):
        return node('sSub', node('e', mr(base)), node('sub', mr(suffix)))
    math = OxmlElement('m:oMath')
    math.append(sub('RMSE', 'xy'))
    math.append(mr(' = '))
    hidden = node('degHide'); hidden.set(qn('m:val'), '1')
    contents = node('e')
    contents.append(node('f', node('num', mr('1')), node('den', mr('N'))))
    operator = node('chr'); operator.set(qn('m:val'), '∑')
    limits = node('limLoc'); limits.set(qn('m:val'), 'undOvr')
    expr = node('e', mr('['))
    for index, axis in enumerate(['x', 'y']):
        if index:
            expr.append(mr(' + '))
        term = node('e', sub('e', axis), mr('(t)'))
        expr.append(node('sSup', term, node('sup', mr('2'))))
    expr.append(mr(']'))
    contents.append(node('nary', node('naryPr', operator, limits),
                         node('sub', mr('t = 1')), node('sup', mr('N')), expr))
    math.append(node('rad', node('radPr', hidden), node('deg'), contents))
    para._p.append(math)

# Page 1: conclusion and conditions.
p('ADAPT 선속도 입력 노이즈와\n지연에 따른 보행 성능 평가', style='Title')
p('actor와 observer가 공유하는 오차에 대한 평면 속도 추종 평가', style='Subtitle')
p('평가일 2026년 10월 21일   |   고정된 Stage 2 정책   |   로컬 RTX 4080 SUPER', style='Caption')
h('1 평가 목적과 핵심 결과')
p('선속도 측정에 노이즈와 지연이 생겼을 때 ADAPT의 보행 성능이 얼마나 달라지는지 평가했다. '
  'actor와 운동량 기반 observer에 동일한 측정 오차를 전달하고, 전진과 좌우 오차를 합친 평면 속도 벡터 RMSE를 대표 지표로 사용했다.')
p('노이즈 ±0.05 m/s에서 평면 속도 RMSE는 무지연 0.0794 m/s, 40 ms 지연 0.0799 m/s로 거의 같았다. '
  '전진 축만 보았을 때의 개선은 좌우 오차 증가로 상쇄됐다. 120 ms 지연에서는 0.1426 m/s로 무지연 대비 약 79.5% 증가했다.', bold=True)
p('총 750회의 30초 시험에서 낙상은 없었다. 그러나 선속도 무오차·무지연 기준에서도 평균 절대 방향 드리프트가 약 40°였다. '
  '따라서 낙상 없이 걷는 능력과 직선 경로를 유지하는 능력은 별도로 판단해야 한다.')
h('2 실험 조건')
table(['항목', '설정'], [
    ['정책', 'Stage 2 model_34998.pt 한 개를 고정하여 평가'],
    ['환경', 'G1 평지 보행  외부 하중 없음'],
    ['시간 간격', '물리 1 ms  정책 및 선속도 측정 20 ms'],
    ['명령', '몸체 기준 vx = 0.8 m/s  vy = 0  yaw rate = 0'],
    ['노이즈', 'xyz 각 축에 U(−a, a)\na = 0, 0.05, 0.1, 0.2, 0.5 m/s'],
    ['고정 지연', '0, 20, 40, 80, 120 ms'],
    ['반복', '25조건 × 10회 × 3개 seed = 750회\nseed 42, 142, 242  조건당 30회'],
    ['평가 구간', '시험당 30초  추종 RMSE는 첫 2초 제외'],
], [32, 138])
caption('자료 [1–3]. 모든 수치는 동일한 정책의 반복 실행 결과이며, 독립적으로 학습한 여러 모델의 비교가 아니다.')

# Page 2: measurement model, metric, complete planar matrix.
page()
h('3 선속도 오차를 적용한 방식')
p('매 정책 주기에 몸체 좌표계의 IMU 선속도를 생성하고, 측정 노이즈를 한 번 적용한 뒤 FIFO 버퍼로 지연시켰다. '
  'actor는 이 측정값을 관측 이력에 넣는다. observer는 같은 센서 오차를 세계 좌표계의 root 선속도로 변환한 값을 사용해 '
  '이상 모델의 동역학과 운동량을 계산한다. 센서 위치와 root 원점의 차이는 각속도와 위치 오프셋으로 보정했다.')
p('추가한 노이즈와 지연은 선속도 xyz에 적용했다. 자세와 각속도는 현재 값을 유지했다. '
  'actor 각속도 노이즈는 ±0.2 rad/s, observer 각속도 노이즈는 0으로 고정했다. '
  '공유 관절 위치 노이즈 ±0.01 rad, 관절 속도 노이즈 ±0.5 rad/s, encoder bias ±0.015 rad도 기존 조건을 유지했다.')
h('4 평면 속도 벡터 RMSE')
equation('RMSE_xy = √{meanₜ[(vₓ − vₓ_cmd)² + (vᵧ − vᵧ_cmd)²]}')
p('eₓ와 eᵧ는 각각 참값 몸체 속도와 명령의 차이이고, N은 평가 구간의 샘플 수다. 전진·좌우 오차의 제곱합을 시간 평균한 뒤 제곱근을 취하며, '
  '축 개수 2로 나누지 않는다. 속도의 방향도 반영하는 지표로, 속력 크기만 비교하는 오차와 다르다.')
p('기존에 저장한 trial별 RMSE로 RMSE_xy = √(RMSE_x² + RMSE_y²)를 복원했다. '
  '각 trial에서 합친 뒤 조건별 평균과 표준편차를 구했으며, 축별 평균값을 합산하지 않았다. '
  '이번 재집계에는 시뮬레이션을 다시 실행하지 않았다. 수직 속도 z는 지표에 포함하지 않고, yaw rate RMSE는 rad/s 단위로 따로 보고한다.')
caption('표 1  전체 조건의 평면 속도 벡터 RMSE  평균 ± 표준편차  단위 m/s  조건당 30회')
records=[]
for a in meta['noise_half_widths_mps']:
    line=[f'±{a:g}']
    for delay in meta['delays_ms']:
        val=by[a,delay]['vxy_rmse_mps']
        line.append(f"{val['mean']:.4f}\n± {val['std']:.4f}")
    records.append(line)
table(['노이즈\nm/s']+[f'{d} ms' for d in meta['delays_ms']], records, [25,29,29,29,29,29], 9)
p('±는 반복 trial 사이의 표준편차이며 신뢰구간이 아니다. 모든 조건에서 낙상은 0/30회였다. '
  '시험 중 낙상이 발생하면 최초 낙상 이후 데이터는 제외하도록 구현되어 있으므로, 추종 지표와 낙상률을 함께 해석해야 한다.')
caption('자료 [1]. 원자료 750건 모두에서 RMSE_xy² = RMSE_x² + RMSE_y² 관계를 확인했다.')

# Page 3: graphs, delay interpretation, corrected control comparison.
page()
h('5 노이즈와 지연에 따른 성능 변화')
pic = doc.add_picture(str(DATA/'sensitivity_curves.png'), width=Mm(170))
doc.paragraphs[-1].paragraph_format.line_spacing = 1.0
pic._inline.docPr.set('descr', '노이즈 5수준별 지연 증가에 따른 평면 속도 벡터 RMSE와 yaw rate RMSE. 오차 막대는 trial 간 표준편차.')
caption('그림 1  actor와 observer가 공유하는 선속도 오차에 대한 민감도  평균 ± 1 SD  자료 [1]')
p('노이즈가 커지거나 지연이 길어지면 전반적으로 평면 속도 추종 오차가 증가했다. '
  '가장 큰 입력 오차인 ±0.5 m/s와 120 ms에서 평면 RMSE는 '
  f"{by[.5,120]['vxy_rmse_mps']['mean']:.4f} m/s였다. "
  '작은 노이즈 조건에서는 20 ms에서 평균 오차가 약간 낮았지만, 80–120 ms에서는 뚜렷하게 커졌다.')
doc.add_heading('짧은 지연의 개선을 해석할 때', 2)
table(['지표', '0 ms', '20 ms', '40 ms'], [
    ['평면 속도 RMSE\nm/s']+[f"{by[.05,d]['vxy_rmse_mps']['mean']:.4f}" for d in [0,20,40]],
    ['전진 속도 RMSE\nm/s']+[f"{by[.05,d]['vx_rmse_mps']['mean']:.4f}" for d in [0,20,40]],
    ['좌우 속도 RMSE\nm/s']+[f"{by[.05,d]['vy_rmse_mps']['mean']:.4f}" for d in [0,20,40]],
    ['Yaw rate RMSE\nrad/s']+[f"{by[.05,d]['yaw_rate_rmse_radps']['mean']:.4f}" for d in [0,20,40]],
], [56,38,38,38], 9)
caption('표 2  노이즈 ±0.05 m/s에서의 조건별 평균  축별 지표는 원인 분석용으로 유지')
p('40 ms에서 전진 오차는 감소했으나 좌우 오차가 증가하여 평면 RMSE는 무지연과 거의 같았다. '
  '20 ms에서는 평면 RMSE가 약 4.6% 낮았다. 보행 중 속도 진동과 피드백 보정 시점이 달라지는 것이 가능한 설명이지만, '
  '원인을 확인한 결과는 아니다. 순수한 지연 자체가 노이즈를 제거하는 필터는 아니며, 최적 지연이나 통계적 유의성도 입증하지 않았다.')
doc.add_heading('Observer에 동일한 오차를 준 영향', 2)
compare=[]
for a,d in [(.05,0),(.05,120),(.5,120)]:
    vals=[math.hypot(r['vx_rmse_mps'],r['vy_rmse_mps']) for r in old_trials if r['noise_mps']==a and r['delay_ms']==d]
    compare.append([f'±{a:g} m/s  {d} ms',f'{statistics.mean(vals):.4f}',f"{by[a,d]['vxy_rmse_mps']['mean']:.4f}"])
table(['입력 조건', 'actor만 오차 적용', 'actor와 observer 공유'], compare,[60,55,55],9)
caption('표 3  평면 RMSE 평균 비교  단위 m/s  자료 [1, 4]. 앞선 actor 단독 결과는 비교군으로만 사용한다.')

# Page 4: limits and audit/reproduction.
page()
h('6 적용 범위와 남은 문제')
p('낙상 없이 30초를 완료한 결과가 직선 경로 추종을 보장하지는 않는다. 선속도 노이즈 0 및 지연 0 조건에서도 '
  f"평균 횡방향 이탈은 {by[0,0]['lateral_drift_m']['mean']:.2f} m, 평균 절대 방향 드리프트는 "
  f"{by[0,0]['heading_drift_deg']['mean']:.1f}°였다. "
  '명령이 몸체 좌표 기준이므로 로봇이 회전하면서 몸체 전진 속도를 유지할 수 있다. 경로 이탈과 회전 지표를 함께 확인해야 한다.')
p('학습의 물리 간격은 5 ms이고 본 평가는 1 ms이다. 정책 주기는 모두 20 ms지만, '
  '물리 적분 조건의 차이가 결과에 영향을 줄 수 있다. 짧은 지연의 개선 원인을 확인하려면 학습과 같은 5 ms 물리 조건에서도 반복해야 한다.')
p('이번 실험은 독립 균등 노이즈와 고정 지연을 사용한다. 실제 FAST-LIO의 bias, 시간 상관, 지연 변동, '
  '자세 오차, 각속도 오차의 공유 및 외부 하중에 대한 민감도는 포함하지 않는다. '
  '정책은 한 개이며 재학습하지 않았다. 따라서 다른 학습 seed나 실제 로봇에서 같은 결과가 난다고 일반화할 수 없다.')
doc.add_heading('후속 평가 우선순위', 2)
p('대표 지표는 평면 속도 벡터 RMSE로 유지하고, yaw rate RMSE·낙상률·경로 이탈을 함께 보고한다. '
  '물리 5 ms에서 같은 실험을 반복한 뒤, 평면 오차를 평균 속도 편차와 시간 변동으로 분해하여 20 ms 조건의 차이를 확인한다. '
  '이후 여러 속도 명령과 실제 센서 로그에 가까운 오차 모델로 평가 범위를 넓힌다.')
h('7 재현 정보와 자료 위치')
p('작업 경로는 /home/rpmsnu/Documents/ADAPT이다. 아래 경로는 이 디렉터리를 기준으로 한다. '
  '모델과 원자료는 보존되어 있으며, 보고서의 표와 그래프는 저장된 결과에서 생성했다.')
caption('체크포인트  logs/rsl_rl/linvel_noise_005/momo_g1_effort_scale/\n2026-10-20_18-00-25-938473_stage2/model_34998.pt')
caption('SHA256  '+meta['checkpoint_sha256'])
code = [
    'CHECKPOINT="logs/rsl_rl/linvel_noise_005/momo_g1_effort_scale/"',
    'CHECKPOINT+="2026-10-20_18-00-25-938473_stage2/model_34998.pt"',
    '.venv/bin/python -m scripts.evaluate_velocity_sensitivity \\',
    '  --checkpoint "$CHECKPOINT" --target shared --seed 42 \\',
    '  --output logs/evaluation/shared_repeat_seed42',
]
para=doc.add_paragraph()
para.paragraph_format.line_spacing=1.1
for i,line in enumerate(code):
    run=para.add_run(line+('\n' if i<len(code)-1 else ''))
    run.font.name='DejaVu Sans Mono'
    run.font.size=Pt(8)
p('새 output 경로로 seed 42, 142, 242를 각각 실행한다. 세 결과 디렉터리를 '
  'scripts.summarize_velocity_sensitivity에 전달하면 평면 RMSE를 집계한다. '
  '입력 공유, 지연 버퍼, 리셋 및 좌표 변환에 대한 5개 검증을 통과했고, 평가 중에도 실제 센서와 전달 신호를 매 스텝 대조했다.')
sources=[
    '[1] 평면 지표와 그래프  logs/evaluation/velocity_sensitivity_shared_2026-10-21_vector/\nsummary.json, trials.json, metadata.json, sensitivity_curves.png',
    '[2] 원래 공유 오차 평가  logs/evaluation/velocity_sensitivity_shared_2026-10-21_seed{42,142,242}/',
    '[3] 검증 및 실행 기록  logs/evaluation/velocity_sensitivity_shared_2026-10-21_combined/\nverification.log, run_seed42.log, run_seed142.log, run_seed242.log',
    '[4] actor 단독 비교군  logs/evaluation/velocity_sensitivity_2026-10-21_combined/trials.json',
]
for source in sources:
    caption(source)
doc.save(OUT)
print(OUT)
